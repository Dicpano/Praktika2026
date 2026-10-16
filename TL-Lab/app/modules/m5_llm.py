"""Модуль 5. Qwen3-0.6B: генерація тестового запитання за текстом (ФВ-5.1…5.6).

Режими: zero-shot (базовий), few-shot і RAG (без зміни ваг), LoRA (донавчання адаптера).
"""
from __future__ import annotations

import json
import random
import re
from functools import lru_cache
from typing import Callable

from core import config, data, experiment
from core.jobs import reporter

DATA_FILE = config.DATA / "quiz" / "quiz.jsonl"
SYSTEM = ("Ти — помічник викладача. За навчальним текстом склади одне тестове запитання українською мовою "
          "з трьома варіантами відповіді, лише один з яких правильний. Відповідай ТІЛЬКИ JSON у форматі: "
          '{"question": "...", "options": ["...", "...", "..."], "answer": "..."}')
SYSTEM_TF = ("Ти — помічник викладача. За навчальним текстом склади одне твердження українською мовою, яке може бути "
             "правдивим або хибним щодо тексту. Відповідай ТІЛЬКИ JSON у форматі: "
             '{"question": "твердження", "options": ["Правда", "Неправда"], "answer": "Правда або Неправда"}')
QTYPES = {"mc": "3 варіанти відповіді", "tf": "Правда / неправда"}
MODES = {"zero": "Zero-shot (базовий рівень)", "few": "Few-shot (3 приклади в запиті)",
         "rag": "RAG (пошук схожих прикладів)", "lora": "LoRA (донавчений адаптер)"}

Generator = Callable[[list[dict]], str]


def load_data(path=DATA_FILE) -> tuple[list[dict], list[dict]]:
    """Навчальні / тестові приклади (20 % — тест), поділ окремо для кожного типу запитань."""
    rows = data.read_jsonl(path)
    for i, r in enumerate(rows, 1):
        r.setdefault("type", "mc")
        n_opt = 3 if r["type"] == "mc" else 2
        if not {"text", "question", "options", "answer"} <= r.keys() or len(r["options"]) != n_opt:
            raise ValueError(f"Приклад {i}: потрібні поля text, question, options ({n_opt}), answer")
    train, test = [], []
    for qt in QTYPES:
        part = [r for r in rows if r["type"] == qt]
        rnd = random.Random(config.SEED)
        idx = list(range(len(part)))
        rnd.shuffle(idx)
        # тест невеликий: генерація на CPU повільна (6 + 3 тексти × кожен режим)
        n_test = min(max(1, round(len(part) * 0.2)), 6 if qt == "mc" else 3) if part else 0
        test += [part[i] for i in idx[:n_test]]
        train += [part[i] for i in idx[n_test:]]
    return train, test


def target(r: dict) -> str:
    return json.dumps({"question": r["question"], "options": r["options"], "answer": r["answer"]}, ensure_ascii=False)


def messages(text: str, examples: list[dict] | None = None, qtype: str = "mc") -> list[dict]:
    msgs = [{"role": "system", "content": SYSTEM_TF if qtype == "tf" else SYSTEM}]
    for e in examples or []:
        msgs += [{"role": "user", "content": e["text"]}, {"role": "assistant", "content": target(e)}]
    msgs.append({"role": "user", "content": text})
    return msgs


def retrieve(text: str, pool: list[dict], k=3, encode=None) -> list[dict]:
    """RAG: k найсхожіших прикладів з навчального набору (MiniLM, косинусна схожість)."""
    if encode is None:
        from modules.text_encoder import encode
    import numpy as np
    m = encode([p["text"] for p in pool])
    v = encode([text])[0]
    return [pool[i] for i in np.argsort(-(m @ v))[:k]]


def parse(out: str) -> dict | None:
    """Витягує перший JSON-об'єкт і перевіряє структуру."""
    out = re.sub(r"<think>.*?</think>", "", out, flags=re.S)
    m = re.search(r"\{.*\}", out, flags=re.S)
    if not m:
        return None
    try:
        d = json.loads(m.group(0))
    except json.JSONDecodeError:
        return None
    if not isinstance(d, dict) or not isinstance(d.get("options"), list) or "question" not in d or "answer" not in d:
        return None
    return d


def valid(p: dict | None, qtype: str = "mc") -> bool:
    if not p:
        return False
    n = 3 if qtype == "mc" else 2
    return len(p["options"]) == n and p["answer"] in p["options"]


def score(outputs: list[str], qtypes: list[str] | None = None) -> dict:
    parsed = [parse(o) for o in outputs]
    qtypes = qtypes or ["mc"] * len(outputs)
    ok = [p for p in parsed if p]
    good = [p for p, q in zip(parsed, qtypes) if valid(p, q)]
    n = max(1, len(outputs))
    return {"valid_json": len(ok) / n, "key_in_options": len(good) / n, "n": len(outputs)}


# ---------------- робота з моделлю (transformers) ----------------

@lru_cache(maxsize=2)
def load_model(with_lora: bool = False):
    import torch
    from transformers import AutoModelForCausalLM, AutoTokenizer
    src = str(config.QWEN_DIR) if (config.QWEN_DIR / "config.json").exists() else config.QWEN_ID
    tok = AutoTokenizer.from_pretrained(src)
    dev = "cuda" if torch.cuda.is_available() else "cpu"
    model = AutoModelForCausalLM.from_pretrained(
        src, torch_dtype=torch.float16 if dev == "cuda" else torch.float32).to(dev)
    if with_lora:
        from peft import PeftModel
        model = PeftModel.from_pretrained(model, str(config.QWEN_LORA_DIR))
    return tok, model.eval()


_POOL = None


def worker_pool():
    """Qwen працює в окремому процесі (spawn): аварія або нестача пам'яті не зупиняють програму."""
    global _POOL
    if _POOL is None:
        import multiprocessing as mp
        from concurrent.futures import ProcessPoolExecutor
        _POOL = ProcessPoolExecutor(max_workers=1, mp_context=mp.get_context("spawn"))
    return _POOL


def reset_worker() -> None:
    """Зупиняє процес Qwen (звільняє пам'ять; після донавчання LoRA модель завантажиться заново)."""
    global _POOL
    if _POOL is not None:
        _POOL.shutdown(wait=False, cancel_futures=True)
        _POOL = None


CRASH_MSG = ("Процес генерації Qwen аварійно завершився — найчастіше через нестачу оперативної пам'яті "
             "(модель займає ≈ 2,5 ГБ). Програма працює далі. Закрийте зайві програми або вкладки браузера "
             "і спробуйте ще раз. Технічні деталі — у файлі app/logs/qwen_worker.log.")


def hf_generator(with_lora=False, max_new_tokens=200) -> Generator:
    from modules import qwen_worker

    def gen(msgs: list[dict]) -> str:
        from concurrent.futures.process import BrokenProcessPool
        try:
            return worker_pool().submit(qwen_worker.generate, str(config.QWEN_DIR), config.QWEN_ID,
                                        str(config.QWEN_LORA_DIR), with_lora, msgs, max_new_tokens).result(timeout=1800)
        except BrokenProcessPool:
            reset_worker()
            raise RuntimeError(CRASH_MSG) from None
    return gen


def lora_available() -> bool:
    return (config.QWEN_LORA_DIR / "adapter_config.json").exists()


def build_messages(mode: str, text: str, train: list[dict], encode=None, qtype: str = "mc") -> list[dict]:
    pool = [r for r in train if r.get("type", "mc") == qtype] or train
    if mode == "few":
        return messages(text, pool[:3], qtype)
    if mode == "rag":
        return messages(text, retrieve(text, pool, 3, encode), qtype)
    return messages(text, None, qtype)  # zero і lora — однаковий запит, різні ваги


def demo(text: str, modes=("zero", "few", "rag"), qtype: str = "mc", job=None) -> dict:
    """Демонстрація: одне запитання за текстом у кожному режимі (фонова задача, модель — в окремому процесі)."""
    j = reporter(job)
    train, _ = load_data()
    outs = {}
    for i, md in enumerate(modes):
        j.set(i / len(modes), f"{MODES[md]}: генерація…")
        if i == 0:
            j.log("перший запуск: завантаження Qwen3-0.6B у пам'ять (≈ 20–60 с), далі — швидше")
        outs[md] = hf_generator(with_lora=md == "lora")(build_messages(md, text, train, qtype=qtype))
        j.log(f"[{md}] готово")
    j.ok("згенеровано у всіх режимах")
    return {"demo": True, "text": text, "qtype": qtype, "outputs": outs}


def make_test(texts: list[str], mode: str = "rag", qtype: str = "mc", job=None) -> dict:
    """Генерує тест: по одному запитанню на кожен текст. Невалідні відповіді позначаються для ручного редагування."""
    j = reporter(job)
    train, _ = load_data()
    gen = hf_generator(with_lora=mode == "lora")
    items = []
    for i, t in enumerate(texts):
        j.set(i / max(1, len(texts)), f"запитання {i + 1}/{len(texts)}")
        p = parse(gen(build_messages(mode, t, train, qtype=qtype)))
        items.append({"text": t, "type": qtype, "ok": valid(p, qtype), **(p or {"question": "", "options": [], "answer": ""})})
        j.log(f"{i + 1}/{len(texts)}: {'✓' if items[-1]['ok'] else 'потребує редагування'}")
    j.ok(f"тест готовий: {sum(x['ok'] for x in items)}/{len(items)} валідних запитань")
    return {"test_items": items, "mode": mode, "qtype": qtype}


def to_gift(items: list[dict], title: str = "TL-Lab") -> str:
    """Формат Moodle GIFT: імпорт у банк запитань Moodle."""
    def esc(x: str) -> str:
        return "".join("\\" + ch if ch in "~=#{}:" else ch for ch in str(x))
    out = [f"// {title}: згенеровано TL-Лабораторією, перевірено викладачем", ""]
    for i, it in enumerate(items, 1):
        if not it.get("question"):
            continue
        if it.get("type") == "tf":
            out.append(f"::Q{i}:: {esc(it['question'])} {{{'TRUE' if it['answer'] == 'Правда' else 'FALSE'}}}")
        else:
            opts = " ".join(("=" if o == it["answer"] else "~") + esc(o) for o in it["options"])
            out.append(f"::Q{i}:: {esc(it['question'])} {{ {opts} }}")
        out.append("")
    return "\n".join(out)


def to_docx(items: list[dict], title: str = "Тест") -> bytes:
    import io

    from docx import Document
    from docx.shared import Pt
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(12)
    doc.add_heading(title, 1)
    keys = []
    for i, it in enumerate([x for x in items if x.get("question")], 1):
        doc.add_paragraph(f"{i}. {it['question']}")
        for n, o in enumerate(it["options"]):
            doc.add_paragraph(f"{'абв'[n]}) {o}", style="List Bullet")
        keys.append(f"{i} — {'абв'[it['options'].index(it['answer'])] if it['answer'] in it['options'] else '?'}")
    doc.add_page_break()
    doc.add_heading("Ключ відповідей", 2)
    doc.add_paragraph("; ".join(keys))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def run_experiment(modes=("zero", "few", "rag", "lora"), gen_factory=None, encode=None, job=None) -> dict:
    j = reporter(job)
    gen_factory = gen_factory or hf_generator
    train, test = load_data()
    modes = [m for m in modes if m != "lora" or lora_available()]
    res, outs = {}, {}
    j.log(f"Qwen3-0.6B · enable_thinking=False · жадібне декодування · тестових текстів: {len(test)}")
    total, done = len(modes) * len(test), 0
    for mode in modes:
        gen = gen_factory(with_lora=(mode == "lora"))
        outs[mode] = []
        for i, r in enumerate(test):
            outs[mode].append(gen(build_messages(mode, r["text"], train, encode, r.get("type", "mc"))))
            done += 1
            j.set(done / total, f"{MODES[mode]}: {i + 1}/{len(test)}")
            j.log(f"[{mode}] згенеровано {i + 1}/{len(test)}")
        res[mode] = score(outs[mode], [r.get("type", "mc") for r in test])
        j.ok(f"[{mode}] валідний JSON {res[mode]['valid_json']:.0%} · ключ серед варіантів {res[mode]['key_in_options']:.0%}")
    best = "lora" if "lora" in res else ("rag" if "rag" in res else "few")
    metrics = {"baseline": res.get("zero", {}), "adapted": res.get(best, {}), "by_mode": res}
    experiment.log_result("m5", {"modes": list(res), "adapted_mode": best, "test": len(test)}, metrics,
                          notes="few-shot і RAG не змінюють ваг; донавчання — лише LoRA")
    out = {"metrics": metrics, "outputs": outs, "test": test}
    experiment.save_extra("m5", out)
    return out


def train_lora(epochs=3, lr=2e-4, r=8, alpha=16, job=None) -> list[float]:
    """Донавчання LoRA на 24 прикладах. Втрата рахується лише на відповіді асистента."""
    import torch
    from peft import LoraConfig, get_peft_model
    experiment.set_seed()
    tok, base = load_model(False)
    base.train()
    model = get_peft_model(base, LoraConfig(r=r, lora_alpha=alpha, lora_dropout=0.05,
                                            target_modules=["q_proj", "v_proj"], task_type="CAUSAL_LM"))
    train, _ = load_data()
    j = reporter(job)
    j.log(f"LoRA r={r}, alpha={alpha}, q_proj/v_proj · lr={lr} · епох {epochs} · прикладів {len(train)}")
    opt = torch.optim.AdamW([p for p in model.parameters() if p.requires_grad], lr=lr)
    hist = []
    for ep in range(epochs):
        random.Random(config.SEED + ep).shuffle(train)
        tot = 0.0
        for n, r_ in enumerate(train):
            j.set((ep + n / len(train)) / epochs, f"епоха {ep + 1}/{epochs}, приклад {n + 1}/{len(train)}")
            msgs = messages(r_["text"], None, r_.get("type", "mc"))
            prompt = tok.apply_chat_template(msgs, tokenize=False, add_generation_prompt=True, enable_thinking=False)
            full = prompt + target(r_) + tok.eos_token
            p_len = len(tok(prompt)["input_ids"])
            enc = tok(full, return_tensors="pt").to(model.device)
            labels = enc["input_ids"].clone()
            labels[:, :p_len] = -100
            loss = model(**enc, labels=labels).loss
            loss.backward()
            opt.step()
            opt.zero_grad()
            tot += loss.item()
        hist.append(tot / len(train))
        j.log(f"епоха {ep + 1}/{epochs}  loss={hist[-1]:.4f}")
    model.save_pretrained(str(config.QWEN_LORA_DIR))
    load_model.cache_clear()
    reset_worker()
    j.ok(f"адаптер збережено → models/{config.QWEN_LORA_DIR.name}/")
    return hist


if __name__ == "__main__":
    r = run_experiment(modes=("zero", "few", "rag"))
    print(json.dumps(r["metrics"], ensure_ascii=False, indent=1))
