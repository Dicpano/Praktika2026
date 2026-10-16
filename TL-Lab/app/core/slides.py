"""Слайди доповіді (.pptx) у стилі програми.

build(key, …)  — 6 слайдів модуля: задача → модель і дані → як адаптували → результати → чому так вийшло → висновок.
build_all(…)   — загальна доповідь ≈ 10 слайдів: титул, що таке transfer learning, по слайду на модуль, підсумок, висновки.
До кожного слайда — тези для виступу в нотатках (режим доповідача PowerPoint).
"""
from __future__ import annotations

import io

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
import numpy as np  # noqa: E402
from pptx import Presentation  # noqa: E402
from pptx.dml.color import RGBColor  # noqa: E402
from pptx.enum.shapes import MSO_SHAPE  # noqa: E402
from pptx.enum.text import PP_ALIGN  # noqa: E402
from pptx.util import Emu, Pt  # noqa: E402

from core import experiment, insights  # noqa: E402
from core.registry import CARDS  # noqa: E402
from core.report import compare_rows  # noqa: E402

BG, PANEL, LINE, CYAN, TEXT, MUTED, GREEN, RED, AMBER = ("05070A", "0B1219", "12313A", "00E5FF", "E6FDFF", "7FA9B3",
                                                         "39FF9A", "FF3B5C", "FFB000")
FONT = "Consolas"
APP = "Transfer Learning-лабораторія"
W, H = 12192000, 6858000   # 16:9
M = 457200                 # поле 0,5″
CW = W - 2 * M
ICON = {"ok": ("✔", GREEN), "wr": ("⚠", AMBER), "info": ("→", CYAN), "tip": ("★", "C792FF")}


def clip(s: str, n: int) -> str:
    s = " ".join(str(s or "").split())
    return s if len(s) <= n else s[: n - 1].rsplit(" ", 1)[0] + "…"


def _rgb(h):
    return RGBColor.from_string(h)


def _slide(prs, notes: str = ""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    f = s.background.fill
    f.solid()
    f.fore_color.rgb = _rgb(BG)
    if notes:
        s.notes_slide.notes_text_frame.text = notes
    return s


def _text(slide, x, y, w, h, text, size=14, color=TEXT, bold=False, align=None):
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        r = p.add_run()
        r.text = line
        r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT
        r.font.color.rgb = _rgb(color)
    return tb


def _rich(slide, x, y, w, h, items, size=13, gap=6):
    """Список пунктів: [(значок, колір, текст)]."""
    tb = slide.shapes.add_textbox(Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, (icon, col, text) in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = icon + "  "
        r.font.size, r.font.name, r.font.bold = Pt(size), FONT, True
        r.font.color.rgb = _rgb(col)
        r = p.add_run()
        r.text = text
        r.font.size, r.font.name = Pt(size), FONT
        r.font.color.rgb = _rgb(TEXT)
    return tb


def _box(slide, x, y, w, h, border=LINE, fill=PANEL, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(int(h)))
    s.fill.solid()
    s.fill.fore_color.rgb = _rgb(fill)
    s.line.color.rgb = _rgb(border)
    s.line.width = Pt(1.25)
    s.shadow.inherit = False
    return s


def _header(slide, key: str | None, title: str, sub: str = ""):
    _text(slide, M, 230000, CW, 300000, (f"root@tl-lab:~/modules/{key}$ ./run" if key else "root@tl-lab:~$ ./report") +
          (f"   ·   {sub}" if sub else ""), 11, MUTED)
    _text(slide, M, 470000, CW, 650000, title, 26 if len(title) <= 44 else 22 if len(title) <= 56 else 19, CYAN, True)
    ln = slide.shapes.add_connector(1, Emu(M), Emu(1130000), Emu(W - M), Emu(1130000))
    ln.line.color.rgb = _rgb(LINE)


def _footer(slide, text=f"{APP} · навчальний макет · лише відкриті, синтетичні або знеособлені дані"):
    _text(slide, M, H - 420000, CW, 300000, text, 9, AMBER)


def _png(fig) -> bytes:
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", facecolor=fig.get_facecolor())
    plt.close(fig)
    return buf.getvalue()


def _fig(w=6.2, h=4.0):
    plt.rcParams.update({"font.family": "DejaVu Sans Mono", "text.color": "#b8f3ff", "axes.labelcolor": "#b8f3ff",
                         "xtick.color": "#7fa9b3", "ytick.color": "#7fa9b3", "axes.edgecolor": "#12313a"})
    fig, ax = plt.subplots(figsize=(w, h), dpi=160)
    fig.patch.set_facecolor("#05070a")
    ax.set_facecolor("#0b1219")
    return fig, ax


def _short(names: dict, k: str, n=9) -> str:
    s = names.get(k, k).split(" (")[0].split(" / ")[0]
    return s if len(s) <= n else s[: n - 1] + "."


def _labels(key: str) -> dict:
    try:
        if key == "m1":
            from modules.m1_vision import LABELS_UA
        elif key == "m2":
            from modules.m2_docs import LABELS_SHORT as LABELS_UA
        elif key == "m3":
            from modules.m3_audio import LABELS_UA
        else:
            return {}
        return LABELS_UA
    except Exception:  # noqa: BLE001
        return {}


def chart(key: str, entry: dict, extra) -> bytes | None:
    """Головний графік модуля, намальований matplotlib у кольорах програми."""
    mt = entry["metrics"]
    fig, ax = _fig()
    if key in ("m1", "m2", "m3"):
        from matplotlib.colors import LinearSegmentedColormap
        cm = np.array(mt["adapted"]["confusion"])
        names = _labels(key)
        labels = [_short(names, x) for x in mt["adapted"]["labels"]]
        ax.imshow(cm, cmap=LinearSegmentedColormap.from_list("c", ["#05070a", "#063a44", "#00e5ff"]))
        fs = 13 if len(labels) <= 6 else 10
        for i in range(cm.shape[0]):
            for jj in range(cm.shape[1]):
                ax.text(jj, i, cm[i, jj], ha="center", va="center", color="#e6fdff", fontsize=fs)
        ax.set_xticks(range(len(labels)), labels, rotation=35, ha="right", fontsize=8)
        ax.set_yticks(range(len(labels)), labels, fontsize=8)
        ax.set_xlabel("прогноз")
        ax.set_ylabel("істина")
        ax.set_title("Матриця помилок адаптованої моделі", color="#e6fdff", fontsize=11)
    elif key == "m4" and extra and extra.get("silhouette_k"):
        ks, vs = zip(*sorted(extra["silhouette_k"].items()))
        ax.plot(ks, vs, "-o", color="#00e5ff")
        ax.set_xlabel("k (кількість кластерів)")
        ax.set_ylabel("silhouette")
        ax.set_title("Вибір кількості кластерів", color="#e6fdff", fontsize=11)
    elif key == "m5" and mt.get("by_mode"):
        by = mt["by_mode"]
        x = np.arange(len(by))
        ax.bar(x - 0.2, [v["valid_json"] for v in by.values()], 0.4, color="#1c6f80", label="валідний JSON")
        ax.bar(x + 0.2, [v["key_in_options"] for v in by.values()], 0.4, color="#00e5ff", label="ключ серед варіантів")
        ax.set_xticks(x, list(by))
        ax.set_ylim(0, 1.1)
        ax.legend(facecolor="#0b1219", edgecolor="#12313a", fontsize=8)
        ax.set_title("Режими генерації", color="#e6fdff", fontsize=11)
    elif key == "m6" and "original_env_after" in mt:
        names = ["змінене\nдо", "змінене\nпісля", "стандарт.\nдо", "стандарт.\nпісля"]
        keys = ["baseline", "adapted", "original_env_before", "original_env_after"]
        ax.bar(names, [mt[k]["mean_reward"] for k in keys], yerr=[mt[k].get("std_reward", 0) for k in keys],
               color=["#ff3b5c", "#00e5ff", "#1c4a55", "#39ff9a"], ecolor="#b8f3ff")
        env = (entry.get("params") or {}).get("env", "cartpole")
        ax.set_ylim(*((-520, 0) if env == "acrobot" else (0, 520)))
        ax.set_ylabel("середня винагорода")
        ax.set_title("До / після донавчання і контроль забування", color="#e6fdff", fontsize=11)
    else:
        plt.close(fig)
        return None
    return _png(fig)


def data_chart(summary: dict) -> bytes | None:
    rows = summary.get("rows") or []
    if not rows:
        return None
    fig, ax = _fig(5.6, 3.9)
    names = [clip(r[0], 22) for r in rows][::-1]
    vals = [r[1] for r in rows][::-1]
    bars = ax.barh(names, vals, color=["#00e5ff" if v >= 10 else "#ffb000" if v >= 4 else "#ff3b5c" for v in vals])
    for b, v in zip(bars, vals):
        ax.text(b.get_width() + max(vals) * 0.02, b.get_y() + b.get_height() / 2, str(v), va="center", fontsize=9,
                color="#e6fdff")
    ax.set_xlim(0, max(vals) * 1.18 + 1)
    ax.set_xlabel(summary.get("unit", ""))
    ax.tick_params(axis="y", labelsize=8)
    ax.set_title("Скільки даних у кожному класі", color="#e6fdff", fontsize=11)
    return _png(fig)


def _kpis(slide, entry, y, x0=M, width=CW, n_max=3):
    rows = [r for r in compare_rows(entry["metrics"]) if not r[0].startswith("СКВ")][:n_max]
    if not rows:
        return

    def nice(v):
        try:
            f = float(v)
        except ValueError:
            return v
        return f"{f:.0f}" if abs(f) >= 10 else f"{f:.2f}".replace(".", ",")
    gap = 150000
    w = (width - gap * (len(rows) - 1)) / len(rows)
    for i, (name, b, a) in enumerate(rows):
        x = x0 + i * (w + gap)
        _box(slide, x, y, w, 1000000, CYAN)
        _text(slide, x + 140000, y + 50000, w - 200000, 300000, name.upper(), 11, MUTED)
        try:
            better = float(a) >= float(b)
        except ValueError:
            better = True
        _text(slide, x + 140000, y + 330000, w - 200000, 600000, f"{nice(b)} → {nice(a)}", 24, GREEN if better else RED, True)


def _flow(slide, a: dict, y):
    """Схема: вхід → заморожена модель → що навчили → результат."""
    ws = [2000000, 3500000, 3300000, 2000000]
    gap = (CW - sum(ws)) / 3
    x = M
    spec = [(a["input"], LINE, PANEL, TEXT), (a["frozen"], "3A5560", "0E1A22", MUTED), (a["trained"], CYAN, "06303A", TEXT),
            (a["output"], LINE, PANEL, TEXT)]
    tags = ["ВХІД", "ГОТОВА МОДЕЛЬ · ЗАМОРОЖЕНО ❄", "НАВЧАЄТЬСЯ 🔥", "РЕЗУЛЬТАТ"]
    for i, ((txt, border, fill, col), w) in enumerate(zip(spec, ws)):
        _box(slide, x, y, w, 1250000, border, fill)
        _text(slide, x + 100000, y + 60000, w - 150000, 260000, tags[i], 9, CYAN if i == 2 else MUTED, True)
        _text(slide, x + 100000, y + 360000, w - 150000, 860000, txt, 12, col, i == 2)
        if i < 3:
            ar = slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW, Emu(int(x + w + 30000)), Emu(int(y + 520000)),
                                        Emu(int(gap - 60000)), Emu(220000))
            ar.fill.solid()
            ar.fill.fore_color.rgb = _rgb(CYAN)
            ar.line.fill.background()
        x += w + gap


def _why_items(why: list[dict], n=6, size=200) -> list[tuple[str, str, str]]:
    order = {"ok": 0, "wr": 1, "info": 2, "tip": 3}
    pick = sorted(why, key=lambda w: order.get(w["t"], 2))[:n]
    pick.sort(key=lambda w: why.index(w))
    return [(ICON.get(w["t"], ("→", CYAN))[0], ICON.get(w["t"], ("→", CYAN))[1], clip(insights.plain(w["md"]), size))
            for w in pick]


def _freeze_table(slide, f, x, y, w):
    if not f or not f.get("rows"):
        return
    rows = f["rows"]
    tbl = slide.shapes.add_table(len(rows) + 1, 3, Emu(int(x)), Emu(int(y)), Emu(int(w)), Emu(300000 * (len(rows) + 1))).table
    head = ["Варіант", "Навчуваних ваг", "Результат"]
    for j, t in enumerate(head):
        c = tbl.cell(0, j)
        c.text = t
    for i, r in enumerate(rows, 1):
        vals = [r["variant"], f"{r['trainable']:,}".replace(",", " ") + f" ({100 * r['trainable'] / max(1, r['total']):.1f} %)",
                f"{r['score']:.0f}" if f.get("metric") == "mean_reward" else f"{r['score']:.3f}"]
        for j, t in enumerate(vals):
            tbl.cell(i, j).text = t
    for i in range(len(rows) + 1):
        for j in range(3):
            c = tbl.cell(i, j)
            c.fill.solid()
            c.fill.fore_color.rgb = _rgb("0E1A22" if i else "06303A")
            for p in c.text_frame.paragraphs:
                for rr in p.runs:
                    rr.font.size, rr.font.name = Pt(10), FONT
                    rr.font.color.rgb = _rgb(CYAN if i == 0 else TEXT)


# ======================================================================== слайди модуля
def build(key: str, entry: dict | None, extra, errors: str, limits: str, conclusion: str, repo: str,
          speaker: str = "") -> bytes:
    c = CARDS[key]
    title = c.title.split(". ", 1)[-1]
    code = f"MOD-{key[1:].zfill(2)}"
    why = insights.why_for(key, entry, extra or {}) if entry else []
    auto = insights.auto_text(key, entry, why)
    errors, limits, conclusion = errors or auto["errors"], limits or auto["limits"], conclusion or auto["conclusion"]
    a = insights.ADAPT[key]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)

    # 1 ---- задача
    need = insights.theory_text(key, "Навіщо це військовим")
    ex = insights.theory_text(key, "Простий приклад")
    s = _slide(prs, f"Модуль {key[1:]}: {title}.\n\nНавіщо: {clip(need, 600)}\n\nПростий приклад: {clip(ex, 500)}\n\n"
                    "Наголосіть: остаточне рішення завжди за людиною, модель лише робить перший прохід.")
    _header(s, key, f"{code} // {title}", speaker)
    _box(s, M, 1350000, CW / 2 - 100000, 4700000)
    _text(s, M + 160000, 1420000, CW / 2 - 400000, 300000, "НАВІЩО ЦЕ ВІЙСЬКОВИМ", 12, CYAN, True)
    _text(s, M + 160000, 1800000, CW / 2 - 400000, 4200000, clip(need, 520), 16)
    x2 = M + CW / 2 + 100000
    _box(s, x2, 1350000, CW / 2 - 100000, 4700000)
    _text(s, x2 + 160000, 1420000, CW / 2 - 400000, 300000, "ПРОСТИЙ ПРИКЛАД", 12, CYAN, True)
    _text(s, x2 + 160000, 1800000, CW / 2 - 400000, 4200000, clip(ex, 520), 16)
    _footer(s)

    # 2 ---- модель і дані
    ds = insights.data_summary(key)
    total = sum(r[1] for r in ds["rows"]) if ds["rows"] else 0
    s = _slide(prs, f"Готова модель: {c.model}. Навчена на: {c.pretrained_on}. Параметрів: {c.params}.\n\n"
                    f"Наших даних: {total} {ds['unit']}. {ds['note']}\n\n"
                    "Ключова думка: модель уже багато знає, тому нам вистачає кількох десятків прикладів на клас, "
                    "а не тисяч.")
    _header(s, key, "Модель і дані", code)
    rows = [("Готова модель", c.model), ("Навчена на", c.pretrained_on), ("Параметрів", c.params),
            ("Ліцензія", c.license), ("Базовий рівень", c.baseline)]
    y = 1350000
    for lbl, val in rows:
        _box(s, M, y, 5900000, 820000)
        _text(s, M + 120000, y + 40000, 5600000, 260000, lbl.upper(), 10, MUTED)
        _text(s, M + 120000, y + 290000, 5600000, 520000, clip(val, 110), 13)
        y += 900000
    png = data_chart(ds)
    if png:
        s.shapes.add_picture(io.BytesIO(png), Emu(M + 6100000), Emu(1350000), width=Emu(CW - 6100000))
    elif key == "m6":
        from modules.m6_rl import ENVS
        _rich(s, M + 6100000, 1350000, CW - 6100000, 3900000,
              [("▸", CYAN, f"{v['title']}: {clip(v.get('about', ''), 150)}") for v in ENVS.values()] +
              [("▸", CYAN, "Зміни умов: фізика (довжина, маса), «вітер» — випадкові поштовхи, «шум датчиків» — "
                           "агент бачить стан з похибкою.")], 13, 10)
    _text(s, M + 6100000, 5500000, CW - 6100000, 700000, clip(ds["note"], 200) +
          (f"\nУсього: {total} {ds['unit']}." if total else ""), 11, MUTED)
    _footer(s)

    # 3 ---- як адаптували
    s = _slide(prs, "Як працює адаптація: " + " ".join(a["steps"]) +
               "\n\nПокажіть на схемі: сіра частина — готова, її не чіпаємо; яскрава — те єдине, що ми навчаємо.")
    _header(s, key, "Як адаптували: перенесення навчання", code)
    _flow(s, a, 1450000)
    _rich(s, M, 3000000, CW, 3000000, [(f"{i}.", CYAN, t) for i, t in enumerate(a["steps"], 1)], 14, 8)
    _footer(s)

    # 4 ---- результати
    hl = insights.headline(key, entry)
    s = _slide(prs, (f"Головна цифра: {hl[0]} {hl[1]} → {hl[2]} (базовий рівень → адаптація).\n\n" if hl else "") +
               "Як читати: ліворуч у табло — базовий рівень, праворуч — після адаптації. "
               + ("Матриця помилок: діагональ — правильні відповіді, решта — плутанина. " if key in ("m1", "m2", "m3") else "")
               + (f"\n\nПараметри запуску: {', '.join(f'{k}={v}' for k, v in (entry or {}).get('params', {}).items() if not isinstance(v, (list, dict)))}"
                  if entry else ""))
    _header(s, key, "Результати: базовий рівень → адаптація", code)
    if entry:
        _kpis(s, entry, 1300000)
        png = chart(key, entry, extra)
        if png:
            s.shapes.add_picture(io.BytesIO(png), Emu(M), Emu(2450000), height=Emu(3850000))
        f = (extra or {}).get("freeze")
        if f:
            _text(s, M + 6300000, 2450000, CW - 6300000, 300000, "ЗАМОРОЖЕНА ЧИ РОЗМОРОЖЕНА", 11, CYAN, True)
            _freeze_table(s, f, M + 6300000, 2800000, CW - 6300000)
            _text(s, M + 6300000, 4000000, CW - 6300000, 2000000, clip(insights.plain(next(
                (w["md"] for w in why if "озморож" in w["md"] or "Часткове" in w["md"]), "")), 320), 11)
        else:
            hint = {"m1": "Діагональ — правильно розпізнані фото. Клітинки поза діагоналлю: рядок — що було, стовпчик — "
                          "за що прийняла модель.",
                    "m2": "Діагональ — правильно визначені типи документів; решта — з чим модель плутає.",
                    "m3": "Діагональ — правильно розпізнані звуки; решта — плутанина (наприклад, Шахед ↔ мопед).",
                    "m4": "Найвища точка кривої — найкраща кількість кластерів за silhouette.",
                    "m5": "Темні стовпчики — правильний формат JSON, світлі — правильний ключ серед варіантів.",
                    "m6": "Червоний — агент до донавчання у зміненому середовищі, блакитний — після; "
                          "зелений — контроль: чи не забув стандартні умови."}[key]
            _box(s, M + 6300000, 2450000, CW - 6300000, 1700000)
            _text(s, M + 6450000, 2520000, CW - 6600000, 300000, "ЯК ЧИТАТИ", 11, CYAN, True)
            _text(s, M + 6450000, 2850000, CW - 6600000, 1300000, hint, 12)
    else:
        _text(s, M, 2000000, CW, 800000, "Запусків ще не було — виконайте адаптацію у вкладці модуля.", 18, AMBER)
    _footer(s)

    # 5 ---- чому так вийшло
    items = _why_items(why, 6, 210)
    s = _slide(prs, "Чому так вийшло — пояснення, яке програма склала з цифр запуску:\n\n" +
               "\n".join(f"- {t}" for _, _, t in items) +
               "\n\nНаголосіть: різницю менше випадкового розкиду (±2√(p(1−p)n) прикладів) не можна вважати ні "
               "покращенням, ні погіршенням.")
    _header(s, key, "Чому так вийшло", code)
    if items:
        _rich(s, M, 1350000, CW, 4900000, items, 14 if len(items) <= 4 else 13, 10)
    else:
        _text(s, M, 2000000, CW, 800000, "Пояснення з'являться після першого запуску адаптації.", 18, AMBER)
    _footer(s)

    # 6 ---- висновок
    s = _slide(prs, f"Висновок: {conclusion}\n\nТипові помилки: {errors}\n\nОбмеження: {limits}\n\n"
                    "Завершіть: повторне використання готових моделей дає результат за хвилини на звичайному ноутбуці; "
                    "покращують його більше даних і розморожування верхніх шарів.")
    _header(s, key, "Висновок", code)
    colw = (CW - 300000) / 3
    for i, (name, txt, col) in enumerate((("ТИПОВІ ПОМИЛКИ", errors, AMBER), ("ОБМЕЖЕННЯ", limits, MUTED),
                                          ("ВИСНОВОК", conclusion, GREEN))):
        x = M + i * (colw + 150000)
        _box(s, x, 1350000, colw, 4600000, col if i == 2 else LINE)
        _text(s, x + 140000, 1420000, colw - 250000, 300000, name, 12, col, True)
        _text(s, x + 140000, 1800000, colw - 250000, 4100000, clip(txt, 430), 12)
    _text(s, M, 6050000, CW, 300000, f"Код: {repo}", 11, MUTED)
    _footer(s)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ======================================================================== загальна доповідь
TASKS = {
    "m1": "Розпізнати на фото техніку, артилерію, БПЛА, замасковані позиції",
    "m2": "Автоматично визначати тип службового документа",
    "m3": "Розпізнавати звуки: Шахед, авіація, техніка, сирена",
    "m4": "Групувати потік повідомлень за темами без розмітки",
    "m5": "Складати тестові запитання за навчальним текстом",
    "m6": "Пристосувати готового агента до змінених умов",
}


def build_all(repo: str = insights.REPO, speaker: str = "") -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(W), Emu(H)
    rows = insights.summary_rows()

    # 1 ---- титул
    s = _slide(prs, "Тема: практичне використання репозиторіїв аналітичних моделей. Покажу, як готові моделі ШІ з "
                    "відкритих репозиторіїв за кілька хвилин пристосовуються до військово-прикладних навчальних задач "
                    "на звичайному ноутбуці, без Інтернету і без тисяч прикладів.")
    _text(s, M, 1500000, CW, 400000, "root@tl-lab:~$ ./report --all", 13, MUTED)
    _text(s, M, 2000000, CW, 1300000, "Transfer Learning-лабораторія", 40, CYAN, True)
    _text(s, M, 3100000, CW, 900000, "Повторне використання готових моделей ШІ\nдля військово-прикладних навчальних задач",
          22, TEXT)
    _text(s, M, 4600000, CW, 700000, "НУОУ · заняття 4.13 «Практичне використання репозиторіїв аналітичних моделей»" +
          (f"\n{speaker}" if speaker else ""), 14, MUTED)
    _footer(s)

    # 2 ---- що таке transfer learning
    s = _slide(prs, "Навчити нейромережу з нуля — це мільйони прикладів, тижні обчислень і дорогі відеокарти. "
                    "Готова модель уже «знає» загальне: як виглядають об'єкти, як звучать двигуни, що означають слова. "
                    "Ми беремо її з репозиторію, заморожуємо і доучуємо лише невелику частину на кількох десятках "
                    "наших прикладів. Аналогія: досвідченого зв'язківця не вчать заново радіотехніки — лише новій станції.")
    _header(s, None, "Що таке transfer learning (перенесення навчання)")
    _flow(s, {"input": "Наші дані\n30–100 прикладів", "frozen": "Готова модель з репозиторію\nмільйони ваг · вже «знає» загальне",
              "trained": "Мала нова частина\nтисячі ваг · вчимо на наших даних", "output": "Рішення нашої задачі\nза хвилини"},
          1450000)
    _rich(s, M, 3100000, CW, 3000000, [
        ("✔", GREEN, "Потрібно в десятки–сотні разів менше даних: 20–30 фото на клас замість тисяч."),
        ("✔", GREEN, "Навчання — хвилини на звичайному ноутбуці без відеокарти; працює без Інтернету."),
        ("✔", GREEN, "Чесна перевірка: кожен модуль порівнює результат із базовим рівнем «без адаптації»."),
        ("⚠", AMBER, "Обмеження: модель знає лише те, що бачила; на чужих даних (ніч, тепловізор) потрібна донавчання."),
    ], 15, 10)
    _footer(s)

    # 3–8 ---- модулі
    for r in rows:
        key = r["key"]
        c = CARDS[key]
        e = experiment.latest(key)
        extra = experiment.load_extra(key) if e else None
        why = insights.why_for(key, e, extra or {}) if e else []
        items = _why_items([w for w in why if w["t"] != "tip"], 3, 170)
        s = _slide(prs, f"{TASKS[key]}. Готова модель: {c.model} ({c.params} параметрів, навчена на {c.pretrained_on}). "
                        f"Що адаптовано: {c.adapted}.\n\n" +
                   (f"Результат: {r['metric']} {r['before']} → {r['after']}.\n" if r["done"] else "Запуску ще не було.\n") +
                   "\n".join(f"- {t}" for _, _, t in items))
        _header(s, key, f"MOD-{key[1:].zfill(2)} // {r['title']}")
        _text(s, M, 1250000, CW, 400000, TASKS[key], 16, TEXT, True)
        _box(s, M, 1750000, 5300000, 1350000, CYAN)
        _text(s, M + 150000, 1800000, 5000000, 300000, ("ГОЛОВНА ЦИФРА · " + r["metric"]).upper(), 11, MUTED)
        _text(s, M + 150000, 2150000, 5000000, 800000, f"{r['before']} → {r['after']}" if r["done"] else "ще не запускали",
              30 if r["done"] else 20, GREEN if r["done"] else AMBER, True)
        if not r["done"]:
            _text(s, M + 5550000, 2000000, CW - 5550000, 1500000, "Запустіть адаптацію цього модуля в програмі — "
                  "тоді тут з'являться графік і пояснення.", 14, MUTED)
        _box(s, M, 3250000, 5300000, 1150000)
        _text(s, M + 150000, 3300000, 5000000, 1050000, f"Готова модель: {c.model}\nПараметрів: {c.params}\n"
              f"Навчена на: {clip(c.pretrained_on, 60)}", 12)
        if items:
            _rich(s, M, 4550000, 5300000, 1800000, items, 11, 4)
        png = chart(key, e, extra) if e else None
        if png:
            s.shapes.add_picture(io.BytesIO(png), Emu(M + 5550000), Emu(1750000), width=Emu(CW - 5550000))
        _footer(s)

    # 9 ---- підсумкова таблиця
    s = _slide(prs, "Підсумок: у кожному модулі готова модель після адаптації порівнюється з базовим рівнем. "
                    "Звертайте увагу не лише на цифри, а й на те, скільки даних і часу знадобилось.")
    _header(s, None, "Підсумок: базовий рівень → адаптація")
    tbl = s.shapes.add_table(len(rows) + 1, 4, Emu(M), Emu(1400000), Emu(CW), Emu(450000 * (len(rows) + 1))).table
    for j, (t, w) in enumerate((("Модуль", 3800000), ("Готова модель", 4100000), ("Метрика", 1900000), ("До → після", 1477600))):
        tbl.cell(0, j).text = t
        tbl.columns[j].width = Emu(w)
    for i, r in enumerate(rows, 1):
        for j, t in enumerate((f"{r['key'][1:].zfill(2)} · {r['title']}", clip(r["model"], 46), r["metric"],
                               f"{r['before']} → {r['after']}" if r["done"] else "—")):
            tbl.cell(i, j).text = t
    for i in range(len(rows) + 1):
        for j in range(4):
            cc = tbl.cell(i, j)
            cc.fill.solid()
            cc.fill.fore_color.rgb = _rgb("0E1A22" if i else "06303A")
            for p in cc.text_frame.paragraphs:
                for rr in p.runs:
                    rr.font.size, rr.font.name = Pt(12), FONT
                    rr.font.bold = j == 3 and i > 0
                    rr.font.color.rgb = _rgb(CYAN if i == 0 else (GREEN if j == 3 else TEXT))
    _footer(s)

    # 10 ---- висновки
    s = _slide(prs, "Висновки. Підкресліть: це навчальний макет на відкритих і синтетичних даних, не для бойового "
                    "застосування; рішення завжди ухвалює людина. Код і посібник — у відкритому репозиторії.")
    _header(s, None, "Висновки")
    _rich(s, M, 1400000, CW, 4300000, [
        ("✔", GREEN, "Готові моделі з відкритих репозиторіїв (PyTorch, Hugging Face, TensorFlow Hub) швидко "
                     "пристосовуються до наших задач: фото, документи, звуки, повідомлення, тести, керування."),
        ("✔", GREEN, "Достатньо десятків прикладів і хвилин на ноутбуці — замість тисяч прикладів і тижнів навчання."),
        ("✔", GREEN, "Кожен результат перевіряється чесно: базовий рівень, тестова вибірка, контроль випадковості й забування."),
        ("⚠", AMBER, "Обмеження: мало даних (особливо записів Шахедів), синтетичні тексти, спрощені симулятори."),
        ("★", "C792FF", "Що далі: більше реальних знеособлених даних, розморожування верхніх шарів, оцінка фахівцями."),
    ], 16, 12)
    _text(s, M, 6000000, CW, 300000, f"Код і посібник: {repo}", 12, CYAN)
    _footer(s)

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
